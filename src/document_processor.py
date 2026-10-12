"""
Módulo de Procesamiento e Ingestión de Documentos Jurídicos (document_processor.py).
Convierte archivos PDF y Word (.docx) a Markdown estructurado (.md) para análisis contextual y RAG.
"""

import io
import re
from pathlib import Path
from typing import Dict, Any, List, Tuple
from pypdf import PdfReader
from docx import Document
from src.config import settings


class DocumentProcessor:
    """Extrae, sanitiza y convierte documentos legales a Markdown limpio."""

    def __init__(self, upload_dir: Path = settings.UPLOAD_DIR):
        self.upload_dir = upload_dir
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def process_file(self, file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """
        Procesa el archivo binario según su extensión y genera una representación en Markdown.
        """
        ext = filename.lower().split(".")[-1]

        if ext == "pdf":
            markdown_content, page_count = self._pdf_to_markdown(file_bytes, filename)
            doc_type = "PDF Legal"
            meta_count = page_count
        elif ext in {"docx", "doc"}:
            markdown_content, para_count = self._docx_to_markdown(file_bytes, filename)
            doc_type = "Word DOCX"
            meta_count = para_count
        elif ext in {"txt", "md"}:
            text = file_bytes.decode("utf-8", errors="ignore")
            markdown_content = self._plain_to_markdown(text, filename)
            doc_type = "Texto Plano / Markdown"
            meta_count = len(text.splitlines())
        else:
            raise ValueError(f"Formato no soportado '.{ext}'. Formatos válidos: PDF (.pdf), Word (.docx), Texto (.txt, .md).")

        # Guardar copia procesada en Markdown
        safe_filename = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", filename)
        md_filename = f"{Path(safe_filename).stem}.md"
        saved_md_path = self.upload_dir / md_filename

        with open(saved_md_path, "w", encoding="utf-8") as f:
            f.write(markdown_content)

        # Segmentar cláusulas / secciones para RAG dinámico
        sections = self._split_into_sections(markdown_content)

        return {
            "filename": filename,
            "tipo_documento": doc_type,
            "archivo_md": md_filename,
            "ruta_guardado": str(saved_md_path),
            "total_elementos": meta_count,
            "caracteres": len(markdown_content),
            "palabras": len(markdown_content.split()),
            "secciones_detectadas": len(sections),
            "secciones": sections,
            "markdown": markdown_content
        }

    def _pdf_to_markdown(self, file_bytes: bytes, filename: str) -> Tuple[str, int]:
        """Convierte páginas de PDF a Markdown estructurado con títulos y párrafos."""
        reader = PdfReader(io.BytesIO(file_bytes))
        total_pages = len(reader.pages)
        md_lines = [
            f"# {Path(filename).stem.replace('_', ' ').title()}",
            f"> **Documento Ingerido:** `{filename}` | **Total de Páginas:** {total_pages}\n"
        ]

        for i, page in enumerate(reader.pages, 1):
            raw_text = page.extract_text() or ""
            cleaned = self._clean_legal_text(raw_text)
            if not cleaned.strip():
                continue

            md_lines.append(f"<!-- Página {i} -->")
            # Detección heurística de títulos de artículos o cláusulas
            for line in cleaned.split("\n"):
                line_str = line.strip()
                if not line_str:
                    continue
                # Si parece un encabezado de artículo o cláusula
                if re.match(r"^(art[ií]culo|cl[aá]usula|secci[oó]n|cap[ií]tulo)\s+\d+", line_str, re.IGNORECASE):
                    md_lines.append(f"\n### {line_str}\n")
                elif re.match(r"^(primera|segunda|tercera|cuarta|quinta|sexta|s[eé]ptima|octava|novena|d[eé]cima)\.?\s*[-—:]", line_str, re.IGNORECASE):
                    md_lines.append(f"\n### {line_str}\n")
                else:
                    md_lines.append(line_str)
            md_lines.append("")

        return "\n".join(md_lines).strip(), total_pages

    def _docx_to_markdown(self, file_bytes: bytes, filename: str) -> Tuple[str, int]:
        """Convierte párrafos y tablas de Word a Markdown con tablas GFM."""
        doc = Document(io.BytesIO(file_bytes))
        md_lines = [
            f"# {Path(filename).stem.replace('_', ' ').title()}",
            f"> **Documento Ingerido:** `{filename}` | **Formato:** Microsoft Word DOCX\n"
        ]

        # Procesar párrafos
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower()
            if "heading 1" in style_name:
                md_lines.append(f"\n## {text}\n")
            elif "heading 2" in style_name or "heading 3" in style_name:
                md_lines.append(f"\n### {text}\n")
            elif re.match(r"^(art[ií]culo|cl[aá]usula|declaraci[oó]n)\s+\d+", text, re.IGNORECASE):
                md_lines.append(f"\n### {text}\n")
            else:
                md_lines.append(text)

        # Procesar tablas en formato Markdown
        if doc.tables:
            md_lines.append("\n## Tablas del Documento\n")
            for t_idx, table in enumerate(doc.tables, 1):
                md_lines.append(f"\n*Tabla {t_idx}:*")
                for r_idx, row in enumerate(table.rows):
                    row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    md_lines.append("| " + " | ".join(row_cells) + " |")
                    if r_idx == 0:
                        md_lines.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
                md_lines.append("")

        return "\n".join(md_lines).strip(), len(doc.paragraphs)

    def _plain_to_markdown(self, text: str, filename: str) -> str:
        """Estructura texto plano a Markdown con encabezado inicial."""
        return f"# {Path(filename).stem.replace('_', ' ').title()}\n\n{text.strip()}"

    def _clean_legal_text(self, text: str) -> str:
        """Elimina ruido tipográfico, pies de página huérfanos y saltos múltiples."""
        t = text.replace("\r\n", "\n").replace("\t", " ")
        t = re.sub(r"P[aá]gina\s+\d+\s+de\s+\d+", "", t, flags=re.IGNORECASE)
        t = re.sub(r"_{3,}", "", t)
        t = re.sub(r"[ ]{2,}", " ", t)
        t = re.sub(r"\n{3,}", "\n\n", t)
        return t

    def _split_into_sections(self, markdown_text: str) -> List[Dict[str, str]]:
        """Divide el Markdown en secciones semánticas para indexación vectorial."""
        # Dividir por encabezados ### o ##
        chunks = re.split(r"\n(?=#{2,3}\s+)", markdown_text)
        sections = []

        for idx, ch in enumerate(chunks):
            ch_str = ch.strip()
            if not ch_str or len(ch_str) < 30:
                continue

            first_line = ch_str.split("\n")[0].replace("#", "").strip()
            title = first_line[:80] if first_line else f"Cláusula {idx+1}"
            sections.append({
                "titulo": title,
                "contenido": ch_str
            })

        return sections


# Instancia singleton para uso en FastAPI
document_processor = DocumentProcessor()
